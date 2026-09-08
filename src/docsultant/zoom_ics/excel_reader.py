"""Read the "Schedule" and "Zoom" sheets of the source workbook into DTOs."""

from __future__ import annotations

from datetime import datetime, time
from pathlib import Path

import openpyxl

from docsultant.zoom_ics.models import WEEKDAYS, ScheduleEntry, ZoomEntry

SCHEDULE_SHEET = "Schedule"
ZOOM_SHEET = "Zoom"


class WorkbookFormatError(ValueError):
    """Raised when the workbook is missing an expected sheet or column."""


def _coerce_time(value: object, *, row_num: int, column: str) -> time:
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%H:%M").time()
        except ValueError as exc:
            raise WorkbookFormatError(
                f"Row {row_num}: cannot parse '{column}' time value {value!r}"
            ) from exc
    raise WorkbookFormatError(f"Row {row_num}: unexpected '{column}' value {value!r}")


def read_schedule(workbook: openpyxl.Workbook) -> list[ScheduleEntry]:
    """Read the Schedule sheet.

    Non-schedule rows (blank rows, the legend block at the bottom of the
    sheet) are recognised because their "Day" cell isn't one of the known
    weekday abbreviations, and are skipped.
    """
    if SCHEDULE_SHEET not in workbook.sheetnames:
        raise WorkbookFormatError(f"Workbook is missing the '{SCHEDULE_SHEET}' sheet")

    sheet = workbook[SCHEDULE_SHEET]
    entries: list[ScheduleEntry] = []

    for row_num, row in enumerate(sheet.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
        day, lesson, type_, active, start_time, end_time, teacher, calendar = row

        if not isinstance(day, str) or day.strip() not in WEEKDAYS:
            continue

        entries.append(
            ScheduleEntry(
                day=day.strip(),
                lesson=(lesson or "").strip(),
                type=(type_ or "").strip(),
                active=(active or "").strip(),
                start_time=_coerce_time(start_time, row_num=row_num, column="Start Time"),
                end_time=_coerce_time(end_time, row_num=row_num, column="End Time"),
                teacher=(teacher or "").strip(),
                calendar=(calendar or "").strip(),
            )
        )

    return entries


def read_zoom_directory(workbook: openpyxl.Workbook) -> dict[str, ZoomEntry]:
    """Read the Zoom sheet into a mapping of teacher name -> :class:`ZoomEntry`."""
    if ZOOM_SHEET not in workbook.sheetnames:
        raise WorkbookFormatError(f"Workbook is missing the '{ZOOM_SHEET}' sheet")

    sheet = workbook[ZOOM_SHEET]
    directory: dict[str, ZoomEntry] = {}

    for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
        teacher, description = row
        if not teacher or not description:
            continue
        teacher = teacher.strip()
        directory[teacher] = ZoomEntry(teacher=teacher, description=description)

    return directory


def load_workbook_data(path: Path) -> tuple[list[ScheduleEntry], dict[str, ZoomEntry]]:
    """Load both sheets of the workbook at ``path``."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return read_schedule(workbook), read_zoom_directory(workbook)
    finally:
        workbook.close()
